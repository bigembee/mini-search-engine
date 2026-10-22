"""Import the workbook's ``Intern Profiles`` sheet into normalized tables."""

from datetime import datetime, time
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from django.core.management.base import BaseCommand, CommandError
from django.utils.text import slugify

from search_app.models import InternProfile, InternSkill, Location, Skill


def excel_time(value):
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, (float, int)):
        seconds = round(float(value) * 24 * 60 * 60)
        return time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
    raw = str(value)
    try:
        return excel_time(float(raw))
    except ValueError:
        return datetime.strptime(raw, "%H:%M:%S").time()


def workbook_rows(path):
    """Yield rows without making the importer depend on a spreadsheet GUI/library."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        with ZipFile(path) as archive:
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = ["".join(node.text or "" for node in item.iter(f"{{{ns}}}t"))
                      for item in shared_root.findall(f"{{{ns}}}si")]
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            workbook_rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relationships = {item.get("Id"): item.get("Target") for item in workbook_rels}
            target = next(
                item.get(f"{{{rel_ns}}}id") for item in workbook.find(f"{{{ns}}}sheets")
                if item.get("name") == "Intern Profiles"
                for _ in [None]
            )
            sheet_path = relationships[target]
            if not sheet_path.startswith("xl/"):
                sheet_path = "xl/" + sheet_path
            root = ElementTree.fromstring(archive.read(sheet_path))
            for xml_row in root.findall(f".//{{{ns}}}sheetData/{{{ns}}}row"):
                values = {}
                for cell in xml_row.findall(f"{{{ns}}}c"):
                    value = cell.find(f"{{{ns}}}v")
                    if value is None:
                        parsed = ""
                    elif cell.get("t") == "s":
                        parsed = shared[int(value.text)]
                    else:
                        parsed = value.text
                    column = "".join(filter(str.isalpha, cell.get("r", "")))
                    values[column] = parsed
                yield [values.get(column, "") for column in
                       ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S")]
    else:
        sheet = load_workbook(path, read_only=True, data_only=True)["Intern Profiles"]
        yield from sheet.iter_rows(values_only=True)


class Command(BaseCommand):
    help = "Import intern profiles from the supplied XLSX workbook."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)

    def handle(self, *args, **options):
        workbook_path = options["workbook"]
        if not workbook_path.exists():
            raise CommandError(f"Workbook not found: {workbook_path}")

        rows = workbook_rows(workbook_path)
        headers = [str(value).strip() for value in next(rows)]
        imported = 0
        for values in rows:
            row = dict(zip(headers, values))
            if not row.get("Intern ID"):
                continue
            location, _ = Location.objects.get_or_create(
                city=row["City"], state=row["State"], country=row["Country"]
            )
            profile, _ = InternProfile.objects.update_or_create(
                intern_id=row["Intern ID"],
                defaults={
                    "full_name": row["Full Name"], "first_name": row["First Name"],
                    "last_name": row.get("Last Name") or "", "role": row["Role / Specialization"],
                    "work_mode": row["Work Mode"].lower(), "work_days": row["Work Days"],
                    "shift_start": excel_time(row["Shift Start"]), "shift_end": excel_time(row["Shift End"]),
                    "daily_hours": row["Daily Hours"], "location": location,
                    "email": row.get("Email") or "", "phone": row.get("Phone") or "",
                    "employment_type": row["Employment Type"].lower().replace(" ", "-"),
                    "status": row["Status"].lower().replace(" ", "-"),
                },
            )
            profile.skills.clear()
            for raw_skill in (row.get("Primary Skills") or "").split(","):
                name = raw_skill.strip()
                if not name:
                    continue
                skill, _ = Skill.objects.get_or_create(name=name, defaults={"slug": slugify(name)})
                InternSkill.objects.get_or_create(intern=profile, skill=skill)
            profile.rebuild_search_keywords()
            profile.save(update_fields=["search_keywords", "updated_at"])
            imported += 1
        self.stdout.write(self.style.SUCCESS(f"Imported {imported} intern profiles."))
